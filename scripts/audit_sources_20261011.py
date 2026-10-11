"""只读扫描 Git 文件、旧能力表和来源摘要，生成审计清单；不更改旧工作簿。"""
import collections
import hashlib
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "evidence" / "audit-20261011"
OUT.mkdir(parents=True, exist_ok=True)
files = subprocess.check_output(["git", "ls-files", "-z"], cwd=ROOT).decode().split("\0")
inventory = []
for name in filter(None, files):
    p = ROOT / name
    if p.is_file():
        inventory.append({"path": name, "bytes": p.stat().st_size,
                          "sha256": hashlib.sha256(p.read_bytes()).hexdigest()})
wbpath = ROOT.parents[1] / "deliverables/engineering-assurance/capability-completion-audit-2026-10-07.xlsx"
wb = openpyxl.load_workbook(wbpath, read_only=True, data_only=False)
sheets = []
for sheet in wb:
    rows = [[v for v in row] for row in sheet.iter_rows(values_only=True)]
    sheets.append({"name": sheet.title, "rows": sheet.max_row, "cols": sheet.max_column,
                   "first_rows": rows[:8], "string_counts": collections.Counter(
                       v for row in rows for v in row if isinstance(v, str) and len(v) < 35)})
result = {"scope": "Git tracked files hashed; this does not mean every path was manually reviewed",
          "head": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT).decode().strip(),
          "file_count": len(inventory), "extensions": collections.Counter(
              Path(r["path"]).suffix for r in inventory), "sheets": sheets,
          "workbook": str(wbpath)}
(OUT / "inventory.json").write_text(json.dumps(inventory, ensure_ascii=False, indent=2), encoding="utf-8")
(OUT / "source-summary.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str),
                                         encoding="utf-8")
print(json.dumps(result, ensure_ascii=False, indent=2, default=str))
