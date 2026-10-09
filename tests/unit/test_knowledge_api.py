"""W3 知识库 · HTTP + 工具通道端到端单测。

覆盖任务书验收清单的后端部分：

* 拖入中文 Markdown → ``status=ready`` → ``POST /api/kb/search`` 命中并带来源；
* 删除文档级联干净（``GET /api/kb/documents`` 为空 + chunks 零残留 + 审计有记录）；
* 适配器卡在未配置时如实显示「未接入」，未接入不返回假列表；
* 未鉴权 401、缺 CSRF 403；
* 损坏文件走「200 + failed + error」而不是静默成功或假成功；
* ``kb.search`` 已进 P1-05 注册中心，``invoke`` 返回**真实**检索结果；
* Chat 调试 SSE 链路：绑定 kb.search 后模型回合出现 ``tool_call`` / ``tool_result``
  且结果来自真实入库文档（工具通道零改动即生效的验证）。
"""

from __future__ import annotations

import asyncio
import json
from typing import Iterator

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import find_yourself.db.kb_models  # noqa: F401
import find_yourself.db.models  # noqa: F401
from find_yourself.api.app import create_app
from find_yourself.config import Settings
from find_yourself.db.base import Base
from find_yourself.db.kb_models import KBChunk
from find_yourself.db.models import AuditEvent
from find_yourself.db.types import TZDateTime
from find_yourself.runtime.gateway import MockModelProvider, ModelGateway
from find_yourself.services.actor import Actor
from find_yourself.services.audit import AuditService
from find_yourself.services.chat_orchestration import ChatOrchestrationService
from find_yourself.services.knowledge import (
    KnowledgeService,
    register_kb_search_tool,
    set_kb_session_factory,
)
from find_yourself.services.knowledge.ingest import KnowledgeIngestService
from find_yourself.services.tool_registry import ToolRegistryService

LOCAL_TOKEN = "dev-token-secret-w3"

MD = """# 数码小屋设计

## 家具布置
家具必须遵守碰撞规则，地板贴图按房间尺寸生成。

## 探险玩法
背景探险包含采集玩法与日常任务三类内容。
"""


# --- 测试基座（沿用 tests/unit/test_preview_sources.py 的 TZ shim 惯例） ------- #
def _tz_result_value(self, value, dialect):  # type: ignore[no-untyped-def]
    if value is not None and value.tzinfo is None:
        return value.replace(tzinfo=__import__("datetime").timezone.utc)
    return value


TZDateTime.process_result_value = _tz_result_value  # type: ignore[method-assign]


def _force_loopback(asgi_app):  # local dev-token 要求 loopback 对端
    async def wrapper(scope, receive, send):
        if scope.get("type") in ("http", "websocket"):
            scope["client"] = ("127.0.0.1", 12345)
        return await asgi_app(scope, receive, send)

    return wrapper


@pytest.fixture()
def session_maker():
    eng = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool, future=True
    )
    Base.metadata.create_all(eng)
    return sessionmaker(bind=eng, expire_on_commit=False, future=True)


@pytest.fixture()
def app(session_maker) -> FastAPI:
    settings = Settings(
        environment="test",
        session_secret="test-session-secret-that-is-long-enough-123456",
        database_url="sqlite://",
        local_token=LOCAL_TOKEN,
        public_url="http://127.0.0.1:8000",
        owner_id="owner",
    )
    application = create_app(session_maker=session_maker, settings=settings)
    set_kb_session_factory(session_maker)
    return application


@pytest.fixture()
def client(app: FastAPI) -> Iterator[TestClient]:
    with TestClient(_force_loopback(app)) as c:
        yield c


@pytest.fixture()
def headers(client: TestClient) -> dict[str, str]:
    r = client.post("/auth/local/dev-token", json={"token": LOCAL_TOKEN})
    assert r.status_code == 200, r.text
    return {"X-CSRF-Token": r.json()["csrf_token"]}


def _upload(client: TestClient, headers: dict[str, str], name: str, data: bytes):
    return client.post(
        f"/api/kb/documents?name={name}",
        content=data,
        headers={**headers, "Content-Type": "application/octet-stream"},
    )


# --- 导入 → 检索 ------------------------------------------------------------- #

def test_upload_markdown_then_search_returns_source(client, headers):
    r = _upload(client, headers, "小屋设计.md", MD.encode("utf-8"))
    assert r.status_code == 200, r.text
    doc = r.json()["document"]
    assert doc["status"] == "ready"
    assert doc["chunk_count"] >= 1
    assert doc["error"] == ""

    s = client.post("/api/kb/search", json={"query": "碰撞规则", "top_k": 5})
    assert s.status_code == 200, s.text
    body = s.json()
    assert body["count"] >= 1
    hit = body["results"][0]
    assert hit["doc_name"] == "小屋设计.md"
    assert hit["source"] == "local"
    assert hit["score"] > 0


def test_documents_endpoint_reports_limits(client, headers):
    _upload(client, headers, "a.md", b"# A\n\n\ufffd\ufffd")
    body = client.get("/api/kb/documents").json()
    assert body["limits"]["max_bytes"] == 20 * 1024 * 1024
    assert ".pdf" in body["limits"]["extensions"]
    assert body["documents"][0]["name"] == "a.md"


def test_unsupported_extension_is_rejected_before_row_created(client, headers, session_maker):
    r = _upload(client, headers, "photo.png", b"\x89PNG\r\n")
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "unsupported_extension"
    session = session_maker()
    assert session.execute(select(KBChunk)).scalars().all() == []
    session.close()


def test_real_docx_upload_becomes_ready_and_is_searchable(client, headers):
    """真实二进制文档闭环：python-docx 生成 DOCX → 上传 → ready → 检索命中。"""
    import io

    import docx

    document = docx.Document()
    document.add_heading("家具布置规范", level=1)
    document.add_paragraph("家具必须遵守碰撞规则，地板贴图按房间尺寸生成。")
    buffer = io.BytesIO()
    document.save(buffer)

    r = _upload(client, headers, "家具规范.docx", buffer.getvalue())
    assert r.status_code == 200, r.text
    doc = r.json()["document"]
    assert doc["status"] == "ready", doc["error"]
    assert doc["chunk_count"] >= 1

    hits = client.post("/api/kb/search", json={"query": "碰撞规则"}).json()
    assert hits["count"] >= 1
    assert hits["results"][0]["doc_name"] == "家具规范.docx"
    assert hits["results"][0]["source"] == "local"


def test_broken_file_returns_200_with_failed_status_and_reason(client, headers):
    r = _upload(client, headers, "broken.pdf", b"%PDF-1.4 garbage")
    assert r.status_code == 200
    doc = r.json()["document"]
    assert doc["status"] == "failed"
    assert "pdf_parse_failed" in doc["error"]


def test_delete_document_cascades_and_audits(client, headers, session_maker):
    doc = _upload(client, headers, "小屋设计.md", MD.encode("utf-8")).json()["document"]
    d = client.delete(f"/api/kb/documents/{doc['id']}", headers=headers)
    assert d.status_code == 200, d.text
    assert d.json()["deleted_chunks"] == doc["chunk_count"]
    assert client.get("/api/kb/documents").json()["documents"] == []

    session = session_maker()
    assert session.execute(select(KBChunk).where(KBChunk.doc_id == doc["id"])).scalars().all() == []
    actions = [e.action for e in session.execute(select(AuditEvent)).scalars().all()]
    assert "kb.document.imported" in actions and "kb.document.deleted" in actions
    session.close()


def test_delete_unknown_document_is_404(client, headers):
    r = client.delete("/api/kb/documents/does-not-exist", headers=headers)
    assert r.status_code == 404
    assert r.json()["error"]["code"] == "kb_document_not_found"


# --- 鉴权 ------------------------------------------------------------------- #

def test_endpoints_require_authentication(client):
    assert client.get("/api/kb/documents").status_code == 401
    assert client.post("/api/kb/search", json={"query": "x"}).status_code == 401
    assert client.get("/api/kb/sources").status_code == 401


def test_upload_requires_csrf_token(client, headers):
    # 已登录（有会话 Cookie）但不带 X-CSRF-Token → 403，而不是放行。
    r = client.post(
        "/api/kb/documents?name=x.md", content=b"# x",
        headers={"Content-Type": "application/octet-stream"},
    )
    assert r.status_code == 403


# --- 适配器（未接入必须诚实） ------------------------------------------------ #

def test_sources_endpoint_reports_ima_and_baidu_not_connected(client, headers):
    from find_yourself.services.knowledge import secret_store

    secret_store.forget("ima")
    body = client.get("/api/kb/sources").json()
    by_id = {s["source_id"]: s for s in body["sources"]}
    # P3 · 备轨上线：/api/kb/sources 现在如实返回三个适配器卡片。
    assert set(by_id) == {"ima", "baidu_pan", "local_files"}
    assert by_id["ima"]["available"] is False
    assert by_id["ima"]["configured"] is False
    assert "未接入" in by_id["ima"]["detail"]
    assert by_id["baidu_pan"]["available"] is False
    assert "未接入" in by_id["baidu_pan"]["detail"]


def test_sync_without_credentials_fails_explicitly(client, headers):
    from find_yourself.services.knowledge import secret_store

    secret_store.forget("ima")
    r = client.post("/api/kb/sources/ima/sync", headers=headers)
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "ima_not_configured"


def test_sync_baidu_pan_reports_not_implemented(client, headers):
    r = client.post("/api/kb/sources/baidu_pan/sync", headers=headers)
    assert r.status_code == 422
    assert r.json()["error"]["code"] == "baidu_pan_not_implemented"


# --- P1-05 工具通道 ---------------------------------------------------------- #

def test_kb_search_tool_is_registered_and_returns_real_results(session_maker):
    set_kb_session_factory(session_maker)
    register_kb_search_tool(tool_registry := ToolRegistryService(persist_dir=None))
    meta = tool_registry.get_tool("kb.search")
    assert meta["entry"]["type"] == "builtin"
    assert meta["parameters"]["required"] == ["query"]
    # owner_id 绝不能是调用方可控参数
    assert "owner_id" not in json.dumps(meta["parameters"])

    session = session_maker()
    KnowledgeIngestService(session, AuditService(session)).ingest_text(
        Actor.owner("owner"), owner_id="owner", name="探险.md",
        text="# 探险\n\n背景探险包含采集玩法与日常任务。\n",
    )
    session.commit()
    session.close()

    receipt = tool_registry.invoke("kb.search", {"query": "采集玩法", "top_k": 3})
    assert receipt["executed"] is True
    assert receipt["result"]["count"] >= 1
    assert receipt["result"]["owner_scope"] == "local_primary"
    assert "采集玩法" in receipt["result"]["results"][0]["content"]


def test_kb_search_tool_reports_no_hit_honestly(session_maker):
    set_kb_session_factory(session_maker)
    set_kb_session_factory(session_maker)
    register_kb_search_tool(reg := ToolRegistryService(persist_dir=None))
    session = session_maker()
    KnowledgeIngestService(session, AuditService(session)).ingest_text(
        Actor.owner("owner"), owner_id="owner", name="x.md", text="# X\n\n无关内容。\n"
    )
    session.commit()
    session.close()
    result = reg.invoke("kb.search", {"query": "量子纠缠退相干时间"}).get("result", {})
    assert result["count"] == 0
    assert result["results"] == []


def test_kb_search_tool_validates_arguments(session_maker):
    set_kb_session_factory(session_maker)
    set_kb_session_factory(session_maker)
    register_kb_search_tool(reg := ToolRegistryService(persist_dir=None))
    with pytest.raises(Exception) as err:
        reg.invoke("kb.search", {"top_k": 3})
    assert "query" in str(err.value)


def test_chat_debug_trace_invokes_kb_search_with_real_results(session_maker):
    """Chat 调试页链路：模型回合 → tool_call(kb.search) → tool_result(真实内容)。"""
    set_kb_session_factory(session_maker)
    register_kb_search_tool(reg := ToolRegistryService(persist_dir=None))
    session = session_maker()
    KnowledgeIngestService(session, AuditService(session)).ingest_text(
        Actor.owner("owner"), owner_id="owner", name="小屋规则.md",
        text="# 小屋规则\n\n家具必须遵守碰撞规则，地板贴图按房间尺寸生成。\n",
    )
    session.commit()
    session.close()

    settings = Settings(
        environment="test",
        session_secret="test-session-secret-that-is-long-enough-123456",
        owner_id="owner",
    )
    marker = json.dumps(
        {"tool_call": {"name": "kb.search", "arguments": {"query": "碰撞规则", "top_k": 3}}},
        ensure_ascii=False,
    )
    provider = MockModelProvider(default_response=marker)
    svc = ChatOrchestrationService(settings, gateway=ModelGateway(provider=provider), registry=reg)

    frames = asyncio.run(_collect(svc.stream_chat(
        Actor.owner("owner"), prompt="p", user_message="用知识库查 碰撞规则",
        model="mock-deterministic", task_id="t-1", tools=[reg.get_tool("kb.search")],
    )))
    events = [json.loads(_data(f)) for f in frames if _event(f) == "tool_result"]
    assert events, [f for f in frames if "tool" in f]
    assert events[0]["name"] == "kb.search"
    assert events[0]["result"]["count"] >= 1
    assert "碰撞规则" in events[0]["result"]["results"][0]["content"]
    # 收尾帧声明本轮真实执行过的工具
    end = [json.loads(_data(f)) for f in frames if _event(f) == "message_end"][0]
    assert end["tools_used"][0]["name"] == "kb.search"
    assert end["tools_used"][0]["executed"] is True


def test_knowledge_service_search_scoped_to_actor_owner(session_maker):
    """跨 owner 隔离（服务层）：owner-1 永远看不到 owner-2 的切片。"""
    session = session_maker()
    svc = KnowledgeService(session, AuditService(session))
    svc.ingest.ingest_text(Actor.owner("owner-2"), owner_id="owner-2",
                           name="secret.md", text="# Secret\n\n口令 rainbow-42。\n")
    session.commit()
    hits = svc.search(Actor.owner("owner-1"), owner_id="owner-1", query="rainbow-42")
    assert hits["count"] == 0
    hits2 = svc.search(Actor.owner("owner-2"), owner_id="owner-2", query="rainbow-42")
    assert hits2["count"] == 1
    session.close()


# --- helpers ---------------------------------------------------------------- #
async def _collect(agen):
    return [frame async for frame in agen]


def _event(frame: str) -> str:
    for line in frame.split("\n"):
        if line.startswith("event: "):
            return line[len("event: ") :]
    return ""


def _data(frame: str) -> str:
    for line in frame.split("\n"):
        if line.startswith("data: "):
            return line[len("data: ") :]
    return "{}"
